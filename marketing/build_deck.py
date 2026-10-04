#!/usr/bin/env python3
"""Build the Keystone marketing deck (PPTX) from KEYSTONE-FEATURES.md content."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ---------------------------------------------------------------- constants
SW, SH = Inches(13.333), Inches(7.5)
NAVY   = RGBColor(0x0E, 0x1B, 0x2A)
BLUE   = RGBColor(0x1B, 0x6A, 0xBF)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)
GREEN  = RGBColor(0x2E, 0x9B, 0x63)
LIGHT  = RGBColor(0xF4, 0xF6, 0xFA)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x1F, 0x29, 0x37)
MUTED  = RGBColor(0x6B, 0x72, 0x80)
FAINT  = RGBColor(0x8B, 0x98, 0xA5)
FONT   = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

SCREENSHOT_DIR = "marketing/screenshots"

# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.25)
        sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sp.shadow.inherit = False
    return sp

def text(s, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(item, (list, tuple)):
            specs = item if item and isinstance(item[0], (list, tuple)) else [item]
        else:
            specs = [(item, {})]
        for spec in specs:
            t = spec[0]
            st = spec[1] if len(spec) > 1 and isinstance(spec[1], dict) else {}
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(st.get("size", size))
            r.font.bold = st.get("bold", bold)
            r.font.color.rgb = st.get("color", color)
    return tb

def bullets(s, x, y, w, h, items, size=13, sub_size=12, line_spacing=1.06):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in items:
        lvl, txt = item if isinstance(item, tuple) and len(item) == 2 and not isinstance(item[1], dict) else (0, item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(5 if lvl == 0 else 3)
        marker = "\u25AA  " if lvl == 0 else "     \u2013  "
        r = p.add_run()
        r.text = marker + txt
        r.font.name = FONT
        r.font.size = Pt(size if lvl == 0 else sub_size)
        r.font.color.rgb = INK if lvl == 0 else MUTED
        r.font.bold = bool(lvl == 0)
    return tb

def header(s, num, kicker, title):
    box(s, 0, 0, SW, Inches(1.05), NAVY)
    box(s, 0, Inches(1.05), SW, Pt(4), AMBER)
    text(s, Inches(0.55), Inches(0.16), Inches(8.5), Inches(0.3),
         [[kicker.upper(), {"size": 11, "bold": True, "color": AMBER}]])
    text(s, Inches(0.55), Inches(0.40), Inches(11.5), Inches(0.62),
         [[title, {"size": 25, "bold": True, "color": WHITE}]])
    text(s, Inches(12.35), Inches(0.30), Inches(0.7), Inches(0.5),
         [[f"{num:02d}", {"size": 18, "bold": True, "color": AMBER}]], align=PP_ALIGN.RIGHT)

def footer(s, page, total):
    text(s, Inches(0.55), Inches(7.12), Inches(6), Inches(0.3),
         [[f"K E Y S T O N E  \u00b7  {page}/{total}", {"size": 9, "color": MUTED}]])
    text(s, Inches(11.3), Inches(7.12), Inches(1.5), Inches(0.3),
         [["v1.2.9", {"size": 9, "color": MUTED}]], align=PP_ALIGN.RIGHT)

def screenshot_placeholder(s, x, y, w, h, title, fname, url):
    box(s, x, y, w, h, LIGHT, line=RGBColor(0xC7, 0xD2, 0xDE), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y + Inches(0.18), w, Inches(0.35),
         [[title, {"size": 12, "bold": True, "color": BLUE}]], align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.55), w, Inches(0.35),
         [[chr(0x2014) + " screenshot goes here " + chr(0x2014), {"size": 10, "color": MUTED}]], align=PP_ALIGN.CENTER)
    text(s, x, y + h - Inches(0.80), w, Inches(0.3),
         [[fname, {"size": 9, "bold": True, "color": INK}]], align=PP_ALIGN.CENTER)
    text(s, x, y + h - Inches(0.52), w, Inches(0.3),
         [[url, {"size": 9, "color": MUTED}]], align=PP_ALIGN.CENTER)
    text(s, x, y + h - Inches(1.08), w, Inches(0.3),
         [[f"save in {SCREENSHOT_DIR}/", {"size": 8, "color": MUTED}]], align=PP_ALIGN.CENTER)

def chip(s, x, y, w, label, fill=AMBER):
    box(s, x, y, w, Inches(0.42), fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y + Inches(0.06), w, Inches(0.32), [[label, {"size": 12, "bold": True, "color": NAVY}]],
         align=PP_ALIGN.CENTER)

# ==================================================== 01 · TITLE
TOTAL = 19
s = slide()
box(s, 0, 0, SW, SH, NAVY)
box(s, 0, Inches(5.02), SW, Pt(4), AMBER)
text(s, Inches(1.0), Inches(1.5), Inches(11.3), Inches(0.4),
     [["REAL ESTATE  \u00b7  WHOLESALING  \u00b7  AI-POWERED", {"size": 13, "bold": True, "color": AMBER}]])
text(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.1),
     [["Keystone", {"size": 66, "bold": True, "color": WHITE}]])
text(s, Inches(1.0), Inches(3.05), Inches(11.3), Inches(1.0),
     [[["The all-in-one CRM for property professionals", {"size": 30, "bold": True, "color": WHITE}]],
      [["Manage motivated-seller leads, deals, buyers, and dispositions \u2014 self-hosted on your own servers.",
        {"size": 16, "color": RGBColor(0xC7, 0xD2, 0xDE)}]]])
chip(s, Inches(1.0), Inches(5.55), Inches(1.9), "Self-hosted")
chip(s, Inches(3.05), Inches(5.55), Inches(1.75), "Multi-tenant")
chip(s, Inches(4.95), Inches(5.55), Inches(2.15), "40+ AI features")
chip(s, Inches(7.25), Inches(5.55), Inches(2.5), "Bring Your Own Key")
chip(s, Inches(9.9), Inches(5.55), Inches(2.6), "GDPR & TCPA ready", fill=GREEN)
text(s, Inches(1.0), Inches(6.55), Inches(11.3), Inches(0.35),
     [["Version 1.2.9 \u00b7 Laravel 12 \u00b7 Open-source", {"size": 12, "color": FAINT}]])

# ==================================================== 02 · TWO PRODUCTS IN ONE
s = slide()
header(s, 2, "Positioning", "One platform. Two businesses.")
cols = [
    ("WHOLESALING", BLUE, [
        "Pipeline: Prospecting \u2192 Dispositions \u2192 Assigned \u2192 Closing",
        "Money field: assignment fee",
        "Property tools: ARV / MAO worksheet, repair costs, distress markers",
        "Modules: Disposition Room, buyer matching, list stacking",
        "Roles: Acquisition Agent, Disposition Agent, Field Scout",
    ]),
    ("REAL ESTATE / BROKERAGE", AMBER, [
        "Pipeline: Listing Agreement \u2192 Active Listing \u2192 Showing \u2192 Closing",
        "Money field: commission",
        "Property tools: list price, listing status, days on market",
        "Modules: Listings, Showings, Open Houses, Leases, A2A contracts",
        "Roles: Listing Agent, Buyers Agent",
    ]),
]
x = Inches(0.55)
for title, accent, items in cols:
    box(s, x, Inches(1.5), Inches(5.95), Inches(4.9), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(s, x, Inches(1.5), Inches(5.95), Inches(0.75), accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.3), Inches(1.63), Inches(5.3), Inches(0.5),
         [[title, {"size": 17, "bold": True, "color": WHITE}]])
    bullets(s, x + Inches(0.35), Inches(2.5), Inches(5.35), Inches(3.8), items, size=13.5)
    x += Inches(6.3)
text(s, Inches(0.55), Inches(6.6), Inches(12.2), Inches(0.35),
     [[["Modules for the other mode are hidden, not missing \u2014 switch modes any time in Settings > Business Mode.",
         {"size": 12, "color": MUTED}]]])
footer(s, 2, TOTAL)

# ==================================================== 03 · LEADS
s = slide()
header(s, 3, "Module", "Lead Management")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Full CRUD with search, sort, and multi-field filtering",
    "14 pipeline statuses and 11 lead sources, all customizable",
    "Dual motivation scoring: system + independent AI score, side by side",
    "Lead Kanban board with drag-and-drop",
    "Quick actions: Call, Email, WhatsApp (auto-normalized number)",
    "Lead photos with lightbox, captions, per-photo delete",
    "CSV import with AI auto-mapping and duplicate handling",
    "Bulk assign / status / delete; saved filter views; keyboard shortcuts",
    (1, "Do Not Contact flag enforced across every outreach action"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Lead Kanban", "04-leads-kanban.png", "/leads/kanban")
footer(s, 3, TOTAL)

# ==================================================== 04 · PIPELINE
s = slide()
header(s, 4, "Module", "Deal Pipeline")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Accordion Kanban: stage rows show count, total value, and fees",
    "Native drag-and-drop between expanded or collapsed stages",
    "Due-diligence countdown with 48-hour urgency alerts",
    "Inline quick-edit: contract price, earnest money, inspection, closing date",
    "Per-deal documents, transaction checklist, and offer management",
    "Automatic activity logging on every stage change",
    "Stage-SLA warnings so nothing stalls",
    (1, "Real estate mode: 11-stage pipeline from lead to closing"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Deal Pipeline", "06-pipeline.png", "/pipeline")
footer(s, 4, TOTAL)

# ==================================================== 05 · DISPOSITION ROOM
s = slide()
header(s, 5, "Module · Wholesale", "Disposition Room")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Opens per deal from the deal detail page",
    "Run all buyer outreach from one screen",
    "Mass-outreach to multiple matched buyers in a single action",
    "Auto deal \u2194 buyer matching when deals reach Dispositions",
    "Transparent match scoring: zip +30, type +25, price +20, state +15",
    "Track per-buyer match status and next actions",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Disposition Room", "08-disposition-room.png", "/disposition/{deal}")
footer(s, 5, TOTAL)

# ==================================================== 06 · PROPERTIES
s = slide()
header(s, 6, "Module", "Properties & Inventory")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Property tracking tied to leads: type, beds, baths, sq ft, year built, lot, condition",
    "Distress markers: tax delinquent, pre-foreclosure, probate, violations, vacant",
    "Real-time financial calculator \u2014 ARV, repairs, MAO, assignment fee",
    "Comparable sales (comps) with the ARV worksheet",
    "Address-normalization for deduplication accuracy",
    "Photo galleries with reorder and primary photo",
    "Inventory / units board with portal-ready export",
    (1, "Field Scout drive-by submission form"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Inventory", "inventory-show.png", "/inventory")
footer(s, 6, TOTAL)

# ==================================================== 07 · BUYERS
s = slide()
header(s, 7, "Module", "Cash Buyers & Roster")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Full buyer database with company info and investment criteria",
    "Reliability scoring that drops when buyers back out",
    "Proof-of-funds (POF) verification with upload and re-score",
    "Buyer transaction log \u2014 their actual track record",
    "Automated deal\u2013buyer matching with explainable scores",
    "One-click notify a matched buyer per deal",
    "CSV import/export and bulk delete",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Buyers", "09-buyers.png", "/buyers")
footer(s, 7, TOTAL)

# ==================================================== 08 · AUTOMATION
s = slide()
header(s, 8, "Automation", "Sequences, Workflows & Campaigns")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Drip sequences: SMS, email, call, voicemail, direct mail with delays",
    "Enrollment progress widget on every lead",
    "Workflow engine: event triggers, step builder, delays, run logs",
    "Pre-built workflow templates for common follow-up flows",
    "Campaign tracker for marketing spend, attribution, and ROI",
    "Email + document templates with merge fields and preview",
    "Outbound webhooks with HMAC signing, retries, auto-disable",
    (1, "Ready for Zapier, Make, and n8n"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Workflow Builder", "13-workflows.png", "/workflows")
footer(s, 8, TOTAL)

# ==================================================== 09 · AI
s = slide()
header(s, 9, "AI Assistant", "40+ AI Features, Bring Your Own Key")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Providers: OpenAI, Claude, Gemini, Ollama, any OpenAI-compatible API",
    "Draft follow-ups across all 7 activity types",
    "AI lead scoring, hot/warm/cold auto-qualification, task suggestions",
    "Deal analysis, stage advice, and offer strategy",
    "Email / subject-line / property / portal description generation",
    "CSV auto-mapping with preview",
    "Weekly digest, pipeline health, DNC risk check, objection library",
    "Briefings on every lead, deal, and buyer",
    (1, "You pay nothing \u2014 tenants use their own API keys"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "AI Drafts", "11-ai-drafts.png", "/leads/{id} \u2192 AI panel")
footer(s, 9, TOTAL)

# ==================================================== 10 · REPORTING
s = slide()
header(s, 10, "Insights", "Reporting & Analytics")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "KPI dashboard with live ApexCharts and AJAX widgets",
    "Pipeline bottleneck detection (days per stage with alerts)",
    "Team leaderboard: leads contacted, offers made, deals closed",
    "Lead-source ROI with per-source cost tracking",
    "Conversion funnel: leads \u2192 contacted \u2192 offer \u2192 contract \u2192 closed",
    "Conversion trend, lead-to-close velocity, top-agent comparison",
    "Goal tracking with forecasting and AI recommendations",
    "CSV exports and print-ready PDF reports",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Reports", "12-reports.png", "/reports")
footer(s, 10, TOTAL)

# ==================================================== 11 · COMMUNICATION & COMPLIANCE
s = slide()
header(s, 11, "Outreach", "Communication & Compliance")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Global activity inbox with type, agent, date, entity filters",
    "Color-coded calendar: tasks, meetings, calls \u2014 click to the lead",
    "iCal feed + Google / Microsoft calendar sync",
    "In-app notification bell with full history",
    "Queued email notifications: assignments, stages, warnings, matches",
    "DNC & TCPA: system-wide stop lists, auto-blocking",
    "Timezone-based contact restrictions (8 AM \u2013 9 PM via ZIP)",
    "GDPR export and anonymized deletion",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Calendar", "10-calendar.png", "/calendar")
footer(s, 11, TOTAL)

# ==================================================== 12 · PORTALS
s = slide()
header(s, 12, "Client-Facing", "Portals & Property Marketing")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Brandable public buyer portal with self-registration",
    "Shared inventory links: clients verify, browse units, flag interest",
    "Interest flags link the unit straight back to the lead record",
    "Embeddable branded web form for lead capture",
    "Listings pushed to Bayut, Dubizzle, and Property Finder",
    "Inbound lead webhooks from property portals",
    (1, "/p/{slug} and /s/{slug} capture prospects without login"),
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Buyer Portal", "buyer-portal.png", "/p/{slug}")
footer(s, 12, TOTAL)

# ==================================================== 13 · REAL ESTATE MODULES
s = slide()
header(s, 13, "Real Estate Mode", "Brokerage Toolkit")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Listings board and mandated-sales pipeline",
    "Inventory / units with portal publishing workflow",
    "Showings with feedback logged to the lead",
    "Open houses with attendee capture",
    "Lease management: create, renew, re-launch search",
    "A2A commission-sharing contracts with external agents",
    "Commission tracking and per-agent plans",
    "CMA tool + portal-readiness checklist",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Listings", "listings.png", "/listings")
footer(s, 13, TOTAL)

# ==================================================== 14 · TEAM & ROLES
s = slide()
header(s, 14, "Team", "Roles & Permissions")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "System roles: Admin, Acquisition, Disposition, Field Scout, Agent",
    "Plus real-estate roles: Listing Agent, Buyers Agent, Cold Call",
    "Custom roles with 33 granular permissions across 9 groups",
    "Team page with manager assignment and hierarchy",
    "Invite by email with welcome notification",
    "Safe member deletion \u2014 records reassigned in one transaction",
    "Impersonation for onboarding and support",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "Settings / Roles", "14-settings.png", "/settings")
footer(s, 14, TOTAL)

# ==================================================== 15 · ADMIN & SECURITY
s = slide()
header(s, 15, "Administration", "Security, Backups & Control")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "Multi-tenant isolation \u2014 every company sees only its own data",
    "Two-factor authentication (TOTP) with enforced 2FA option",
    "Pluggable SSO framework (Google, Azure, Okta via plugins)",
    "Full searchable audit log with before/after values",
    "One-command database backups with scheduled cleanup",
    "Safe updates: ZIP upload, staging, auto-backup, recovery snapshots",
    "Local or S3-compatible cloud storage",
    "System health check and factory-reset safety",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "System / Settings", "14-settings.png", "/settings?tab=system")
footer(s, 15, TOTAL)

# ==================================================== 16 · INTEGRATIONS & API
s = slide()
header(s, 16, "Developer", "Integrations & API")
bullets(s, Inches(0.55), Inches(1.4), Inches(7.3), Inches(5.5), [
    "REST API (v1) for leads, deals, buyers, properties, activities, stats",
    "Per-tenant API keys with rate limiting and request logs",
    "OpenAPI 3.0 spec + interactive docs",
    "Outbound webhooks with HMAC signing and auto-retry",
    "Plugin system: ZIP install, hooks, menus, widgets, settings tabs",
    "SMS gateway abstraction for any provider",
    "Google & Microsoft calendar + photo connections",
])
screenshot_placeholder(s, Inches(8.15), Inches(1.5), Inches(4.6), Inches(5.0),
                      "API Docs", "api-docs.png", "/api-docs")
footer(s, 16, TOTAL)

# ==================================================== 17 · INTERNATIONALIZATION
s = slide()
header(s, 17, "Global", "Built for Every Market")
bullets(s, Inches(0.55), Inches(1.5), Inches(7.3), Inches(4.4), [
    "7 language packs: English, Dutch, German, French, Spanish, Portuguese, Italian",
    "Add any language by dropping in a JSON file",
    "Built-in language editor with search and progress tracking",
    "38 currencies and 50 countries with per-tenant defaults",
    "50+ timezones grouped by region",
    "Imperial / metric units (sq ft \u2194 m\u00b2, acres \u2194 hectares)",
    "Localized views, calculators, CSV exports, and AI prompts",
], size=14)
box(s, Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.15), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(0.9), Inches(5.78), Inches(11.5), Inches(0.35),
     [[["Also included:", {"size": 12, "bold": True, "color": BLUE}]]])
text(s, Inches(0.9), Inches(6.14), Inches(11.5), Inches(0.5),
     [[["Dark mode \u00b7 PWA + offline \u00b7 Mobile-responsive \u00b7 Onboarding wizard \u00b7 Whitelabel branding",
         {"size": 12, "color": INK}]]])
footer(s, 17, TOTAL)

# ==================================================== 18 · TECH STACK
s = slide()
header(s, 18, "Architecture", "Tech Stack")
rows = [
    ("Backend", "Laravel 12 (PHP 8.2+)"),
    ("Frontend", "Tabler Admin (Bootstrap 5) with dark mode"),
    ("Charts", "ApexCharts"),
    ("Kanban", "Native HTML5 drag & drop"),
    ("Database", "MySQL 8.0+ / MariaDB 10.6+"),
    ("AI", "OpenAI \u00b7 Anthropic \u00b7 Gemini \u00b7 Ollama \u00b7 OpenAI-compatible (BYOK)"),
    ("Storage", "Local or S3-compatible"),
    ("API", "REST + per-tenant key auth, OpenAPI spec"),
    ("Deployment", "Self-hosted \u00b7 Docker-ready \u00b7 Web installer \u00b7 Auto-updates"),
]
y = Inches(1.45)
for i, (k, v) in enumerate(rows):
    box(s, Inches(0.55), y, Inches(3.3), Inches(0.52), BLUE if i % 2 == 0 else NAVY,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(0.75), y + Inches(0.09), Inches(3.0), Inches(0.35),
         [[k, {"size": 13, "bold": True, "color": WHITE}]])
    text(s, Inches(4.15), y + Inches(0.09), Inches(8.6), Inches(0.35),
         [[v, {"size": 13, "color": INK}]])
    y += Inches(0.60)
footer(s, 18, TOTAL)

# ==================================================== 19 · CLOSING
s = slide()
box(s, 0, 0, SW, SH, NAVY)
box(s, 0, Inches(4.55), SW, Pt(4), AMBER)
text(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
     [["CONTACT & NEXT STEPS", {"size": 13, "bold": True, "color": AMBER}]])
text(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.0),
     [[["Keystone \u2014 ready to run your deals.", {"size": 40, "bold": True, "color": WHITE}]]])
text(s, Inches(1.0), Inches(3.55), Inches(11.3), Inches(0.8),
     [[["Self-hosted \u00b7 multi-tenant \u00b7 AI-powered \u00b7 compliance-ready",
         {"size": 18, "color": RGBColor(0xC7, 0xD2, 0xDE)}]]])
chip(s, Inches(1.0), Inches(5.0), Inches(3.3), "Try the demo dataset")
chip(s, Inches(4.5), Inches(5.0), Inches(3.1), "Self-host on your server")
chip(s, Inches(7.8), Inches(5.0), Inches(2.9), "Open-source (MIT)", fill=GREEN)
text(s, Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.5),
     [["Load a full demo: php artisan db:seed  \u00b7  login admin@demo.com / password",
        {"size": 12, "color": FAINT}]])

# ---------------------------------------------------------------- save
os.makedirs(SCREENSHOT_DIR, exist_ok=True)
out = "marketing/Keystone-Marketing-Deck.pptx"
prs.save(out)
print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides")